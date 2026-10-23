# -*- coding: utf-8 -*-
"""Gera a planilha CNAE · Descrição · Anexo a partir do Concla.

NAO ALTERA O APP. Apenas *importa* (somente leitura) o motor de enquadramento
do sistema (`cnaes.sugerir_anexo` -> `enquadramento.MOTOR.enquadrar`) para que
o anexo gravado na planilha seja exatamente o que o sistema entende.

Entrada:  CNAE_Subclasses_2_3_Estrutura_Detalhada.xlsx (planilha oficial do
          Concla: Seção > Divisão > Grupo > Classe > Subclasse).
Saída:    planilha PLANA com 3 colunas:

              A = CNAE         (ex.: 0111-3/01)
              B = Descrição    (denominação oficial da subclasse)
              C = Anexo        (anexo do sistema; se houver mais de um
                                possível, lista os candidatos)

Motivo: o sistema enquadra pelo CNAE de 7 dígitos (a subclasse). Linhas de
seção/divisão/grupo/classe não têm 7 dígitos e por isso NÃO são enquadráveis;
portanto só as subclasses entram na planilha.

Uso (PowerShell, dentro da pasta simples_app):

    python preencher_anexo_planilha.py "C:\\caminho\\CNAE_Subclasses_2_3_Estrutura_Detalhada.xlsx"

    # opcional: gravar em outro arquivo (o original nunca e sobrescrito)
    python preencher_anexo_planilha.py planilha.xlsx -o subclasses_anexo.xlsx

Por padrao gera "<nome>_ANEXO.xlsx" na mesma pasta do original.
"""

from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from cnaes import (  # noqa: E402
    apenas_digitos,
    anexos_candidatos,
    formatar_cnae,
    sugerir_anexo,
)

# Colunas da planilha de saida: A = CNAE, B = Descricao, C = Anexo.
COLS_SAIDA = [
    ("CNAE", 14),
    ("Descrição", 90),
    ("Anexo", 22),
]


def _txt(v) -> str:
    if v is None:
        return ""
    return " ".join(str(v).split()).strip()


def _ler_subclasses(ws):
    """Le a estrutura hierarquica do Concla e devolve [(codigo, descricao)].

    Colunas: Seção | Divisão | Grupo | Classe | Subclasse | Denominação.
    Cada linha preenche APENAS um nivel; a descricao da subclasse esta na
    mesma linha da subclasse, na coluna de denominacao.
    """
    linha_cab = None
    col_sub = col_desc = None
    # O Concla nao coloca todos os rotulos na mesma linha: 'Denominação' fica
    # numa linha acima (ex.: linha 3) e 'Subclasse' logo abaixo (linha 4).
    # Atencao: o titulo da linha 1 tambem contem a palavra "Denominação" e a
    # palavra "subclasses", por isso as linhas de titulo precisam ser puladas.
    max_row_cab = min(20, ws.max_row or 20)
    maior_linha = 0
    for linha in range(1, max_row_cab + 1):
        rotulos = {}
        for c in range(1, (ws.max_column or 1) + 1):
            rotulo = _txt(ws.cell(row=linha, column=c).value).lower()
            if rotulo:
                rotulos[rotulo] = c
        # Linha de cabecalho de verdade: tem os rotulos curtos, um em cada
        # coluna (o titulo da linha 1 e uma frase longa).
        if col_sub is None and "subclasse" in rotulos:
            col_sub = rotulos["subclasse"]
            linha_cab = linha
        if col_desc is None:
            for rotulo, c in rotulos.items():
                if len(rotulo) <= 20 and "denomina" in rotulo:
                    col_desc = c
                    break
        if col_sub is not None and col_desc is not None:
            break
        maior_linha = max(maior_linha, linha)
    if linha_cab is None:
        linha_cab = maior_linha

    if linha_cab is None or col_sub is None:
        raise SystemExit("Nao encontrei o cabecalho com 'Subclasse' na planilha.")
    if col_desc is None:
        col_desc = col_sub + 1

    itens = []
    for row in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
        sub = row[col_sub - 1] if len(row) >= col_sub else None
        den = row[col_desc - 1] if len(row) >= col_desc else None
        codigo = apenas_digitos(sub)
        if len(codigo) != 7:
            continue  # secao/divisao/grupo/classe: nao sao subclasses
        itens.append((codigo, _txt(den)))
    return itens


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("planilha", help="Caminho do .xlsx original do Concla")
    ap.add_argument("-o", "--saida", default="", help="Arquivo de saida")
    ap.add_argument("--aba", default="", help="Nome da aba (padrao: a primeira)")
    args = ap.parse_args(argv)

    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        raise SystemExit("Instale o openpyxl: python -m pip install openpyxl")

    origem = os.path.abspath(args.planilha)
    if not os.path.isfile(origem):
        raise SystemExit("Arquivo nao encontrado: %s" % origem)

    destino = args.saida or os.path.splitext(origem)[0] + "_ANEXO.xlsx"
    if os.path.abspath(destino) == origem:
        raise SystemExit(
            "A saida nao pode ser o proprio original. Escolha outro nome em -o."
        )

    wb_in = load_workbook(origem, read_only=True)
    ws_in = wb_in[args.aba] if args.aba else wb_in.worksheets[0]
    itens = _ler_subclasses(ws_in)
    wb_in.close()

    if not itens:
        raise SystemExit("Nenhuma subclasse (7 digitos) encontrada na planilha.")

    # -- Planilha nova, com CNAE, Descricao e Anexo ------------------------
    wb = Workbook()
    ws = wb.active
    ws.title = "CNAE Subclasses"

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    borda = Border(*[Side(style="thin", color="BFBFBF")] * 4)

    for i, (nome, largura) in enumerate(COLS_SAIDA, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borda
        ws.column_dimensions[c.column_letter].width = largura

    ambiguos = 0
    sem_anexo = 0
    for linha, (codigo, descricao) in enumerate(itens, start=2):
        sugestao = sugerir_anexo(codigo, descricao)
        candidatos = anexos_candidatos(codigo, descricao) or []

        if len(candidatos) > 1:
            # Ambiguidade (ex.: depende do Fator "r"): grava a lista, para
            # ficar explicito que a definicao nao e automatica.
            ambiguos += 1
            anexo = " / ".join("Anexo %s" % a for a in candidatos)
            alinhamento_anexo = "left"
        else:
            anexo = sugestao.get("anexo")
            if not anexo and candidatos:
                anexo = candidatos[0]
            anexo = ("Anexo %s" % anexo) if anexo else ""
            alinhamento_anexo = "center"
            if not anexo:
                sem_anexo += 1

        valores = [formatar_cnae(codigo), descricao, anexo]
        for i, valor in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=i, value=valor)
            c.border = borda
            if i == 2:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            elif i == 3:
                c.alignment = Alignment(
                    horizontal=alinhamento_anexo, vertical="top"
                )
            else:
                c.alignment = Alignment(vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:C%d" % (len(itens) + 1)
    wb.save(destino)

    print("Planilha gerada: %s" % destino)
    print("Subclasses processadas: %d" % len(itens))
    print("Com 2+ anexos possiveis (lista gravada): %d" % ambiguos)
    if sem_anexo:
        print("Sem anexo identificado: %d" % sem_anexo)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
