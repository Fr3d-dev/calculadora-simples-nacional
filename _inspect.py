import openpyxl,sys,re
sys.stdout.reconfigure(encoding="utf-8")
a="CNAE_Subclasses_2_3_Estrutura_Detalhada_BAK.xlsx"
ws1=openpyxl.load_workbook(a)["Estrutura Det. CNAE Subclass2.3"]
ws2=openpyxl.load_workbook("CNAE_Subclasses_2_3_Estrutura_Detalhada.xlsx")["Estrutura Det. CNAE Subclass2.3"]
code=re.compile(r"^\s*\d{2}\.\d{2}\s*[–\-]")
tot1=[r for r in range(1,ws1.max_row+1) if ws1.cell(r,6).value and code.match(str(ws1.cell(r,6).value))]
tot2=[r for r in range(1,ws2.max_row+1) if ws2.cell(r,6).value and code.match(str(ws2.cell(r,6).value))]
print("F com item (BAK):",len(tot1)," no arquivo novo:",len(tot2))
print("ultima linha preenchida BAK:",tot1[-1] if tot1 else None)
print("---- ultimas 10 ----")
for r in tot1[-10:]:
    print(r,"| E:",ws1.cell(r,5).value,"| F:",str(ws1.cell(r,6).value)[:90])
print("---- secao R/S no BAK ----")
for r in range(2250,ws1.max_row+1):
    f=ws1.cell(r,6).value
    if f and code.match(str(f)):
        if r<2262 or r>2380: print(r,"|E|",ws1.cell(r,5).value,"|F|",str(f)[:60])
